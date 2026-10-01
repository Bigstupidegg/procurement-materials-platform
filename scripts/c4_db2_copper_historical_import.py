"""Guarded C4-DB2 World Bank Copper historical pilot importer.

This module has no ambient network or persistence behavior.  Callers must
explicitly supply workbook bytes (or call the explicit download helper), an
outside-repository private-data root, a migrated private database, and the two
authority timestamps.  It does not establish PIT, research, backtest,
prediction, or production readiness.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
import io
import os
from pathlib import Path
import tempfile
from typing import Any, Mapping

import openpyxl
import requests

from scripts.c4_private_research_db import (
    REPOSITORY_ROOT,
    PersistenceConflictError,
    PrivateResearchDatabase,
)
from scripts.c4_rd_contract import (
    CALENDAR_ROLES,
    CalendarAssignment,
    Observation,
    ObservationVersion,
    build_real_historical_observation,
    canonical_json_bytes,
    canonical_timestamp,
    raw_payload_hash,
)
from scripts.sync_world_bank import (
    COLUMN_CONFIG_PATH,
    MATERIALS_PATH,
    discover_monthly_url,
    download_xlsx,
    normalize_text,
    parse_period,
    read_json,
    resolve_columns,
)


SOURCE_ID = "WORLD_BANK"
SUBJECT_ID = "copper"
INSTRUMENT_ID = "copper_world_bank_monthly"
ACCESS_CHANNEL = "WORLD_BANK_OFFICIAL_XLSX"
SOURCE_COLUMN = "Copper"
MARKET_OR_VENUE = "WORLD_BANK_PINK_SHEET"
METRIC_ID = "MONTHLY_PRICE"
TRANSFORMATION_VERSION = "C4_DB2_WB_COPPER_MONTHLY_V1@1.0.0"
SOURCE_REGISTRY_VERSION = "C4_DB2_WORLD_BANK_SOURCE_V1@1.0.0"
INSTRUMENT_VERSION = "C4_DB2_WORLD_BANK_COPPER_MONTHLY_V1@1.0.0"
RIGHTS_PROFILE_ID = "WORLD_BANK_PINK_SHEET_DB2_PRIVATE_RESEARCH_V1"
COLLECTOR_VERSION = "C4_DB2_WORLD_BANK_COPPER_COLLECTOR_V1@1.0.0"
PILOT_START = "2022-01"
PILOT_END = "2025-12"
PILOT_MONTHS = 48
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RIGHTS_EVIDENCE_REFERENCE = (
    "Dataset: World Bank Commodity Markets Data / Pink Sheet; "
    "World Bank Data Catalog dataset: 0038238 — Commodity Prices - History and Projections; "
    "Classification: Public; License: Creative Commons Attribution 4.0; "
    "Official evidence families: World Bank Commodity Markets / Using this Data; "
    "World Bank Summary Terms of Use; World Bank Data Access and Licensing; "
    "Human-approved scope: C4-DB2 private historical pilot"
)


class DB2CopperImportError(RuntimeError):
    """Base error for the fail-closed DB2 Copper pilot."""


class DB2SourceFormatError(DB2CopperImportError):
    """Raised when the workbook does not match the approved source shape."""


class PrivateRawStorageError(DB2CopperImportError):
    """Raised when exact raw evidence cannot be stored without overwrite."""


@dataclass(frozen=True, slots=True)
class CopperMonthlyValue:
    period: str
    value: Decimal


@dataclass(frozen=True, slots=True)
class ParsedCopperWorkbook:
    observations: tuple[CopperMonthlyValue, ...]
    source_unit: str
    instrument_active_from: str


@dataclass(frozen=True, slots=True)
class CopperImportResult:
    raw_payload_hash: str
    raw_relative_path: str
    capture_id: str
    observation_content_hashes: tuple[str, ...]
    observation_version_content_hashes: tuple[str, ...]


def _month_sequence(start: str, end: str) -> tuple[str, ...]:
    year, month = (int(part) for part in start.split("-"))
    end_year, end_month = (int(part) for part in end.split("-"))
    result: list[str] = []
    while (year, month) <= (end_year, end_month):
        result.append(f"{year:04d}-{month:02d}")
        if month == 12:
            year, month = year + 1, 1
        else:
            month += 1
    return tuple(result)


EXPECTED_PILOT_PERIODS = _month_sequence(PILOT_START, PILOT_END)


def _decimal_from_cell(value: Any, period: str) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise DB2SourceFormatError(f"Copper value for {period} is not an authoritative numeric cell")
    try:
        number = value if type(value) is Decimal else Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise DB2SourceFormatError(f"Copper value for {period} is invalid") from exc
    if not number.is_finite() or number <= 0:
        raise DB2SourceFormatError(f"Copper value for {period} must be positive and finite")
    return number


def _copper_material(materials: list[dict[str, Any]]) -> dict[str, Any]:
    matches = [
        item for item in materials
        if item.get("id") == SUBJECT_ID and item.get("worldBankColumn") == SOURCE_COLUMN
    ]
    if len(matches) != 1:
        raise DB2SourceFormatError("repository configuration does not identify exactly one Copper subject")
    return matches[0]


def parse_world_bank_copper_workbook(
    raw_xlsx: bytes,
    *,
    materials: list[dict[str, Any]] | None = None,
    column_config: dict[str, Any] | None = None,
) -> ParsedCopperWorkbook:
    """Parse exact authoritative Decimal values for the frozen 48-month pilot."""

    if type(raw_xlsx) is not bytes or not raw_xlsx.startswith(b"PK"):
        raise DB2SourceFormatError("source payload must be exact XLSX bytes")
    materials = read_json(MATERIALS_PATH) if materials is None else materials
    column_config = read_json(COLUMN_CONFIG_PATH) if column_config is None else column_config
    material = _copper_material(materials)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(raw_xlsx), read_only=True, data_only=True, keep_links=False,
        )
    except Exception as exc:
        raise DB2SourceFormatError("source payload is not a readable XLSX workbook") from exc

    try:
        sheet_name = column_config["sheetName"]
        if sheet_name not in workbook.sheetnames:
            raise DB2SourceFormatError(f"required worksheet is missing: {sheet_name}")
        worksheet = workbook[sheet_name]
        header_row = int(column_config["headerRow"])
        unit_row = int(column_config["unitRow"])
        data_start_row = int(column_config["dataStartRow"])
        period_column = int(column_config["periodColumnIndex"])
        header_values = list(next(worksheet.iter_rows(
            min_row=header_row, max_row=header_row, values_only=True,
        )))
        unit_values = list(next(worksheet.iter_rows(
            min_row=unit_row, max_row=unit_row, values_only=True,
        )))
        try:
            copper_column = resolve_columns(header_values, [material], column_config)[SUBJECT_ID]
        except Exception as exc:
            raise DB2SourceFormatError("Copper column is missing or ambiguous") from exc
        source_unit = normalize_text(
            unit_values[copper_column - 1] if len(unit_values) >= copper_column else None
        )
        if source_unit.casefold() != "$/mt":
            raise DB2SourceFormatError("Copper source unit is not the approved $/mt unit")

        all_values: list[CopperMonthlyValue] = []
        for row in worksheet.iter_rows(min_row=data_start_row, values_only=True):
            period = parse_period(row[period_column - 1] if len(row) >= period_column else None)
            if period is None:
                continue
            raw_value = row[copper_column - 1] if len(row) >= copper_column else None
            all_values.append(CopperMonthlyValue(period, _decimal_from_cell(raw_value, period)))
        if not all_values:
            raise DB2SourceFormatError("workbook contains no recognized monthly Copper observations")
        periods = tuple(item.period for item in all_values)
        if len(periods) != len(set(periods)):
            raise DB2SourceFormatError("workbook contains duplicate Copper months")
        if periods != tuple(sorted(periods)):
            raise DB2SourceFormatError("workbook Copper months are not strictly increasing")

        selected = tuple(
            item for item in all_values if PILOT_START <= item.period <= PILOT_END
        )
        selected_periods = tuple(item.period for item in selected)
        if selected_periods != EXPECTED_PILOT_PERIODS or len(selected) != PILOT_MONTHS:
            raise DB2SourceFormatError(
                "Copper pilot must contain exactly 48 consecutive months from 2022-01 through 2025-12"
            )
        return ParsedCopperWorkbook(
            observations=selected,
            source_unit=source_unit,
            instrument_active_from=f"{all_values[0].period}-01",
        )
    except DB2SourceFormatError:
        raise
    except (KeyError, TypeError, ValueError, StopIteration, IndexError) as exc:
        raise DB2SourceFormatError("World Bank workbook/config structure drifted") from exc
    finally:
        workbook.close()


def _private_root(private_data_root: str | Path) -> Path:
    if str(private_data_root).strip() == "":
        raise PrivateRawStorageError("PRIVATE_DATA_ROOT must be explicit")
    root = Path(private_data_root).expanduser().resolve(strict=False)
    repository_root = REPOSITORY_ROOT.resolve(strict=True)
    if root == repository_root or repository_root in root.parents:
        raise PrivateRawStorageError("PRIVATE_DATA_ROOT must be outside the repository")
    if root.exists() and not root.is_dir():
        raise PrivateRawStorageError("PRIVATE_DATA_ROOT must identify a directory")
    return root


def store_raw_world_bank_xlsx(
    raw_xlsx: bytes,
    private_data_root: str | Path,
) -> tuple[str, str, Path]:
    """Create the exact content-addressed raw object without ever overwriting it."""

    if type(raw_xlsx) is not bytes:
        raise PrivateRawStorageError("raw workbook evidence must be exact bytes")
    digest = raw_payload_hash(raw_xlsx)
    root = _private_root(private_data_root)
    relative = Path("raw") / "world-bank" / "pink-sheet" / "copper" / f"{digest}.xlsx"
    target = root / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        if not target.is_file():
            raise PrivateRawStorageError("content-addressed raw path is not a file")
        stored = target.read_bytes()
        if stored != raw_xlsx or raw_payload_hash(stored) != digest:
            raise PrivateRawStorageError("content-addressed raw path conflicts with exact bytes")
        return digest, relative.as_posix(), target

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", dir=target.parent, prefix=f".{digest}.", suffix=".tmp", delete=False,
        ) as temporary:
            temporary.write(raw_xlsx)
            temporary.flush()
            os.fsync(temporary.fileno())
            temporary_path = Path(temporary.name)
        try:
            os.link(temporary_path, target)
        except FileExistsError:
            stored = target.read_bytes()
            if stored != raw_xlsx or raw_payload_hash(stored) != digest:
                raise PrivateRawStorageError("concurrent raw evidence creation conflicted")
        stored = target.read_bytes()
        if stored != raw_xlsx or raw_payload_hash(stored) != digest:
            raise PrivateRawStorageError("stored raw evidence failed exact read-back")
        return digest, relative.as_posix(), target
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def world_bank_copper_calendar_assignments() -> tuple[CalendarAssignment, ...]:
    statuses = {
        "MARKET": "NOT_APPLICABLE",
        "PUBLICATION": "UNVERIFIED",
        "LOCAL_OPERATIONAL": "NOT_APPLICABLE",
        "SCHEDULER": "NOT_APPLICABLE",
    }
    return tuple(
        CalendarAssignment(SUBJECT_ID, role, statuses[role]) for role in CALENDAR_ROLES
    )


def _month_end(period: str) -> str:
    year, month = (int(part) for part in period.split("-"))
    if month == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, month + 1, 1)
    return date.fromordinal(next_month.toordinal() - 1).isoformat()


def build_copper_observation(
    item: CopperMonthlyValue,
    *,
    source_unit: str,
    collected_at: str,
    created_at: str,
) -> Observation:
    semantic_data = {
        "value": item.value,
        "currency": "USD",
        "unit": "USD/MT",
        "source_column": SOURCE_COLUMN,
        "source_unit": source_unit,
        "frequency": "MONTHLY",
    }
    return build_real_historical_observation(
        source_id=SOURCE_ID,
        source_record_identifier=f"WORLD_BANK:PINK_SHEET:COPPER:{item.period}",
        metric_id=METRIC_ID,
        instrument_id=INSTRUMENT_ID,
        source_period_type="MONTH",
        source_market_date=None,
        source_period_start_date=f"{item.period}-01",
        source_period_end_date=_month_end(item.period),
        calendar_assignments=world_bank_copper_calendar_assignments(),
        collected_at=collected_at,
        created_at=created_at,
        semantic_data=semantic_data,
    )


def _stable_version(observation: Observation, digest: str) -> ObservationVersion:
    return ObservationVersion(
        observation_id=observation.observation_id,
        source_version_or_release_key=f"WORLD_BANK_PINK_SHEET_XLSX_SHA256:{digest}",
        stable_version_key=(
            f"{observation.source_record_identifier}:XLSX_SHA256:{digest}"
        ),
        raw_payload_hash=digest,
        transformation_version=TRANSFORMATION_VERSION,
        parent_version_id=None,
        revision_available_at=None,
        collected_at=observation.collected_at,
        observed_at=None,
        created_at=observation.created_at,
        semantic_data=observation.semantic_data,
    )


def _existing(database: PrivateResearchDatabase, loader: str, key: str) -> Mapping[str, Any] | None:
    try:
        return getattr(database, loader)(key)
    except KeyError:
        return None


def _persist_registries(
    database: PrivateResearchDatabase,
    parsed: ParsedCopperWorkbook,
    *,
    created_at: str,
    rights_reviewed_at: str,
) -> None:
    source = _existing(database, "load_source_registry", SOURCE_ID)
    database.persist_source_registry(
        source_id=SOURCE_ID,
        source_name="World Bank Commodity Markets Data (Pink Sheet)",
        access_channel=ACCESS_CHANNEL,
        default_timezone="UTC",
        data_classification="PUBLIC_LICENSED_SOURCE_DATA",
        registry_version=SOURCE_REGISTRY_VERSION,
        active=True,
        created_at=source["created_at"] if source else created_at,
    )
    subject = _existing(database, "load_subject_registry", SUBJECT_ID)
    subject_snapshot = {
        "subject_id": SUBJECT_ID,
        "name": "Copper",
        "scope": "WORLD_BANK_PINK_SHEET_COPPER_ONLY",
    }
    database.persist_subject_registry(
        subject_id=SUBJECT_ID,
        subject_snapshot=subject_snapshot,
        registered_at=subject["registered_at"] if subject else created_at,
    )
    database.persist_instrument_registry(
        instrument_id=INSTRUMENT_ID,
        subject_id=SUBJECT_ID,
        source_id=SOURCE_ID,
        source_symbol=SOURCE_COLUMN,
        market_or_venue=MARKET_OR_VENUE,
        metric_id=METRIC_ID,
        quote_type="PERIOD_AVERAGE",
        term="MONTHLY",
        currency="USD",
        unit="USD/MT",
        source_period_type="MONTH",
        instrument_version=INSTRUMENT_VERSION,
        active_from=parsed.instrument_active_from,
        active_to=None,
    )
    rights = _existing(database, "load_source_usage_rights", RIGHTS_PROFILE_ID)
    database.persist_source_usage_rights(
        rights_profile_id=RIGHTS_PROFILE_ID,
        source_id=SOURCE_ID,
        access_channel=ACCESS_CHANNEL,
        instrument_scope=INSTRUMENT_ID,
        automated_access="ALLOWED",
        private_storage="ALLOWED",
        historical_archive="ALLOWED",
        internal_analysis="ALLOWED",
        backtest="REVIEW_REQUIRED",
        prediction="REVIEW_REQUIRED",
        internal_display="REVIEW_REQUIRED",
        internet_display="REVIEW_REQUIRED",
        redistribution="REVIEW_REQUIRED",
        evidence_reference=RIGHTS_EVIDENCE_REFERENCE,
        review_status="HUMAN_REVIEWED_DB2_PRIVATE_HISTORICAL_PILOT",
        effective_from="2026-10-01",
        effective_to=None,
        reviewed_at=rights["reviewed_at"] if rights else rights_reviewed_at,
    )


def import_world_bank_copper_historical(
    raw_xlsx: bytes,
    *,
    private_data_root: str | Path,
    database: PrivateResearchDatabase,
    collected_at: str,
    created_at: str,
    rights_reviewed_at: str,
    source_locator_safe: str = ACCESS_CHANNEL,
) -> CopperImportResult:
    """Persist the guarded DB2 pilot after a caller's separate Human Gate."""

    if type(database) is not PrivateResearchDatabase:
        raise DB2CopperImportError("database must be an exact PrivateResearchDatabase")
    canonical_rights_reviewed_at = canonical_timestamp(rights_reviewed_at)
    existing_rights = _existing(database, "load_source_usage_rights", RIGHTS_PROFILE_ID)
    if existing_rights is not None and existing_rights["reviewed_at"] is None:
        raise PersistenceConflictError(
            "existing immutable DB2 rights profile has no reviewed_at timestamp"
        )
    parsed = parse_world_bank_copper_workbook(raw_xlsx)
    digest, relative_path, _absolute_path = store_raw_world_bank_xlsx(
        raw_xlsx, private_data_root,
    )
    _persist_registries(
        database,
        parsed,
        created_at=created_at,
        rights_reviewed_at=canonical_rights_reviewed_at,
    )

    raw_metadata = _existing(database, "load_raw_payload", digest)
    database.persist_raw_payload(
        raw_payload_hash=digest,
        relative_path=relative_path,
        content_type=XLSX_CONTENT_TYPE,
        byte_size=len(raw_xlsx),
        data_classification="PUBLIC_LICENSED_SOURCE_DATA",
        first_seen_at=raw_metadata["first_seen_at"] if raw_metadata else collected_at,
    )
    capture_id = f"WORLD_BANK:PINK_SHEET:COPPER:{digest}"
    capture = _existing(database, "load_raw_capture", capture_id)
    database.persist_raw_capture(
        capture_id=capture_id,
        source_id=SOURCE_ID,
        instrument_id=INSTRUMENT_ID,
        access_channel=ACCESS_CHANNEL,
        source_locator_safe=(
            capture["source_locator_safe"] if capture else source_locator_safe
        ),
        collected_at=capture["collected_at"] if capture else collected_at,
        collection_status="SUCCESS",
        raw_payload_hash=digest,
        collector_version=COLLECTOR_VERSION,
        rights_profile_id=RIGHTS_PROFILE_ID,
        error_class=None,
        created_at=capture["created_at"] if capture else created_at,
    )

    observation_hashes: list[str] = []
    version_hashes: list[str] = []
    for item in parsed.observations:
        observation = build_copper_observation(
            item,
            source_unit=parsed.source_unit,
            collected_at=collected_at,
            created_at=created_at,
        )
        version = _stable_version(observation, digest)
        existing_hashes = database.observation_version_content_hashes(
            version.observation_version_id
        )
        if existing_hashes:
            if len(existing_hashes) != 1:
                raise PersistenceConflictError(
                    "existing DB2 observation version has contradictory content snapshots"
                )
            stored_version = database.load_observation_version(existing_hashes[0])
            if (
                stored_version.identity_projection() != version.identity_projection()
                or canonical_json_bytes(stored_version.semantic_data)
                != canonical_json_bytes(version.semantic_data)
                or stored_version.parent_version_id is not None
                or stored_version.revision_available_at is not None
                or stored_version.observed_at is not None
            ):
                raise PersistenceConflictError(
                    "existing DB2 observation version contradicts the exact raw import"
                )
            observation = build_copper_observation(
                item,
                source_unit=parsed.source_unit,
                collected_at=stored_version.collected_at,
                created_at=stored_version.created_at,
            )
            version = stored_version
        observation_hashes.append(database.persist_observation(observation))
        version_hashes.append(database.persist_observation_version(version))

    return CopperImportResult(
        raw_payload_hash=digest,
        raw_relative_path=relative_path,
        capture_id=capture_id,
        observation_content_hashes=tuple(observation_hashes),
        observation_version_content_hashes=tuple(version_hashes),
    )


def import_world_bank_copper_historical_path(
    workbook_path: str | Path,
    *,
    rights_reviewed_at: str,
    **kwargs: Any,
) -> CopperImportResult:
    """Explicit local-file API; the path itself is never persisted."""

    return import_world_bank_copper_historical(
        Path(workbook_path).read_bytes(),
        rights_reviewed_at=rights_reviewed_at,
        **kwargs,
    )


def download_official_world_bank_monthly_xlsx(
    *,
    session: requests.Session | None = None,
) -> tuple[bytes, Mapping[str, Any]]:
    """Explicit opt-in network helper.  Importing this module never calls it."""

    config = read_json(COLUMN_CONFIG_PATH)
    owned_session = session is None
    active_session = requests.Session() if session is None else session
    try:
        url = discover_monthly_url(config, active_session)
        raw, metadata = download_xlsx(url, config, active_session)
        return raw, metadata
    finally:
        if owned_session:
            active_session.close()
