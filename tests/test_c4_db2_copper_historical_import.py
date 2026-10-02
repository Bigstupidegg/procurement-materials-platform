from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

import duckdb
import openpyxl

from scripts.c4_db2_copper_historical_import import (
    ACCESS_CHANNEL,
    EXPECTED_PILOT_PERIODS,
    INSTRUMENT_ID,
    PILOT_MONTHS,
    RIGHTS_EVIDENCE_REFERENCE,
    RIGHTS_PROFILE_ID,
    SOURCE_ID,
    TRANSFORMATION_VERSION,
    DB2SourceFormatError,
    PrivateRawStorageError,
    _decimal_from_cell,
    build_copper_observation,
    import_world_bank_copper_historical,
    parse_world_bank_copper_workbook,
    store_raw_world_bank_xlsx,
)
from scripts.c4_private_research_db import PrivateResearchDatabase, REPOSITORY_ROOT
from scripts.c4_rd_contract import (
    BACKTEST_ENABLED_SOURCES,
    PIT_ENABLED_SOURCES,
    RD3_OPEN_BLOCKERS,
    RESEARCH_ENABLED_SOURCES,
    SAFETY_FLAGS,
    ContractError,
    raw_payload_hash,
)


def workbook_bytes(
    *,
    periods: tuple[str, ...] = EXPECTED_PILOT_PERIODS,
    copper_header: str = "Copper",
    duplicate_copper: bool = False,
    source_unit: str = "($/mt)",
    value_overrides: dict[str, object] | None = None,
) -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Monthly Prices"
    worksheet.cell(5, 1, "Date")
    worksheet.cell(5, 2, "Zinc")
    worksheet.cell(5, 3, copper_header)
    worksheet.cell(6, 3, source_unit)
    if duplicate_copper:
        worksheet.cell(5, 4, "Copper ($/mt)")
        worksheet.cell(6, 4, source_unit)
    overrides = value_overrides or {}
    for offset, period in enumerate(periods, start=7):
        worksheet.cell(offset, 1, period.replace("-", "M"))
        worksheet.cell(offset, 2, 3000 + offset)
        worksheet.cell(offset, 3, overrides.get(period, 8000 + offset / 100))
        if duplicate_copper:
            worksheet.cell(offset, 4, 9000 + offset / 100)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class WorldBankCopperParserTests(unittest.TestCase):
    def test_semantic_copper_column_parses_exact_48_months_as_decimal(self) -> None:
        parsed = parse_world_bank_copper_workbook(
            workbook_bytes(copper_header="Copper ($/mt)")
        )
        self.assertEqual(len(parsed.observations), PILOT_MONTHS)
        self.assertEqual(parsed.observations[0].period, "2022-01")
        self.assertEqual(parsed.observations[-1].period, "2025-12")
        self.assertTrue(all(type(item.value) is Decimal for item in parsed.observations))
        self.assertEqual(parsed.source_unit, "($/mt)")
        self.assertEqual(parsed.instrument_active_from, "2022-01-01")

    def test_legacy_dollar_per_metric_ton_unit_remains_approved(self) -> None:
        parsed = parse_world_bank_copper_workbook(workbook_bytes(source_unit="$/mt"))
        self.assertEqual(parsed.source_unit, "$/mt")

    def test_missing_or_duplicate_copper_column_fails_closed(self) -> None:
        with self.assertRaises(DB2SourceFormatError):
            parse_world_bank_copper_workbook(workbook_bytes(copper_header="Tin"))
        with self.assertRaises(DB2SourceFormatError):
            parse_world_bank_copper_workbook(workbook_bytes(duplicate_copper=True))

    def test_duplicate_missing_and_out_of_order_months_fail_closed(self) -> None:
        duplicate = EXPECTED_PILOT_PERIODS + (EXPECTED_PILOT_PERIODS[-1],)
        missing = EXPECTED_PILOT_PERIODS[:-1]
        out_of_order = list(EXPECTED_PILOT_PERIODS)
        out_of_order[5], out_of_order[6] = out_of_order[6], out_of_order[5]
        for periods in (duplicate, missing, tuple(out_of_order)):
            with self.subTest(periods=len(periods)), self.assertRaises(DB2SourceFormatError):
                parse_world_bank_copper_workbook(workbook_bytes(periods=periods))

    def test_invalid_nonpositive_boolean_text_and_nonfinite_values_fail_closed(self) -> None:
        for value in (0, -1, True, "9123.45"):
            with self.subTest(value=value), self.assertRaises(DB2SourceFormatError):
                parse_world_bank_copper_workbook(
                    workbook_bytes(value_overrides={"2024-01": value})
                )
        for value in (float("nan"), float("inf"), Decimal("NaN")):
            with self.subTest(value=value), self.assertRaises(DB2SourceFormatError):
                _decimal_from_cell(value, "2024-01")

    def test_unknown_source_units_fail_closed(self) -> None:
        for source_unit in ("USD/MT", "$/ton", "($/ton)", "cents/lb", "", "$ / mt"):
            with self.subTest(source_unit=source_unit), self.assertRaises(DB2SourceFormatError):
                parse_world_bank_copper_workbook(workbook_bytes(source_unit=source_unit))

    def test_month_boundaries_and_calendar_assignments_are_exact(self) -> None:
        parsed = parse_world_bank_copper_workbook(workbook_bytes())
        leap = next(item for item in parsed.observations if item.period == "2024-02")
        observation = build_copper_observation(
            leap,
            source_unit=parsed.source_unit,
            collected_at="2026-10-01T00:00:00Z",
            created_at="2026-10-01T00:00:01Z",
        )
        self.assertEqual(observation.source_period_start_date, "2024-02-01")
        self.assertEqual(observation.source_period_end_date, "2024-02-29")
        self.assertEqual(parsed.source_unit, "($/mt)")
        self.assertEqual(observation.semantic_data["source_unit"], "($/mt)")
        self.assertEqual(observation.semantic_data["unit"], "USD/MT")
        self.assertEqual(
            {item.role: item.status for item in observation.calendar_assignments},
            {
                "MARKET": "NOT_APPLICABLE",
                "PUBLICATION": "UNVERIFIED",
                "LOCAL_OPERATIONAL": "NOT_APPLICABLE",
                "SCHEDULER": "NOT_APPLICABLE",
            },
        )
        self.assertTrue(all(item.subject_id == "copper" for item in observation.calendar_assignments))
        self.assertIsNone(observation.source_publication_at)
        self.assertIsNone(observation.source_available_at)
        self.assertIsNone(observation.channel_available_at)
        self.assertIsNone(observation.observed_at)


class WorldBankCopperRawStorageTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory(prefix="c4_db2_raw_")
        self.private_root = Path(self.temporary_directory.name)
        self.raw = workbook_bytes()

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_raw_hash_and_content_addressed_storage_are_exact_and_idempotent(self) -> None:
        first = store_raw_world_bank_xlsx(self.raw, self.private_root)
        second = store_raw_world_bank_xlsx(self.raw, self.private_root)
        self.assertEqual(first, second)
        digest, relative_path, target = first
        self.assertEqual(digest, raw_payload_hash(self.raw))
        self.assertEqual(target.read_bytes(), self.raw)
        self.assertEqual(
            relative_path,
            f"raw/world-bank/pink-sheet/copper/{digest}.xlsx",
        )

    def test_repository_local_and_empty_private_roots_are_rejected(self) -> None:
        for root in ("", REPOSITORY_ROOT, REPOSITORY_ROOT / "private"):
            with self.subTest(root=root), self.assertRaises(PrivateRawStorageError):
                store_raw_world_bank_xlsx(self.raw, root)

    def test_existing_content_addressed_path_with_different_bytes_fails_closed(self) -> None:
        digest = raw_payload_hash(self.raw)
        target = (
            self.private_root / "raw" / "world-bank" / "pink-sheet" / "copper"
            / f"{digest}.xlsx"
        )
        target.parent.mkdir(parents=True)
        target.write_bytes(b"contradictory bytes")
        with self.assertRaises(PrivateRawStorageError):
            store_raw_world_bank_xlsx(self.raw, self.private_root)


class WorldBankCopperImportTests(unittest.TestCase):
    CANONICAL_RIGHTS_REVIEWED_AT = "2026-09-30T16:00:00Z"

    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory(prefix="c4_db2_import_")
        self.private_root = Path(self.temporary_directory.name)
        db_directory = self.private_root / "db"
        db_directory.mkdir()
        self.database_path = db_directory / "pilot.duckdb"
        self.database = PrivateResearchDatabase(self.database_path)
        self.database.initialize_schema(
            applied_at="2026-10-01T00:00:00Z",
            code_commit_sha="SYNTHETIC_DB2_TEST_COMMIT",
        )
        self.raw = workbook_bytes()

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def _import(
        self,
        raw: bytes | None = None,
        *,
        suffix: str = "00",
        rights_reviewed_at: str = "2026-10-01T00:00:00+08:00",
        source_locator_safe: str = ACCESS_CHANNEL,
    ):
        return import_world_bank_copper_historical(
            self.raw if raw is None else raw,
            private_data_root=self.private_root,
            database=self.database,
            collected_at=f"2026-10-01T00:00:{suffix}Z",
            created_at=f"2026-10-01T00:01:{suffix}Z",
            rights_reviewed_at=rights_reviewed_at,
            source_locator_safe=source_locator_safe,
        )

    def test_rights_reviewed_at_is_mandatory_and_invalid_value_has_no_side_effects(self) -> None:
        common = {
            "private_data_root": self.private_root,
            "database": self.database,
            "collected_at": "2026-10-01T00:00:00Z",
            "created_at": "2026-10-01T00:01:00Z",
        }
        with self.assertRaises(TypeError):
            import_world_bank_copper_historical(self.raw, **common)
        with self.assertRaises(ContractError):
            import_world_bank_copper_historical(
                self.raw,
                rights_reviewed_at="2026-10-01",
                **common,
            )
        self.assertFalse((self.private_root / "raw").exists())
        connection = duckdb.connect(str(self.database_path), read_only=True)
        try:
            for table in (
                "source_registry", "subject_registry", "instrument_registry",
                "source_usage_rights", "raw_payload", "raw_capture",
                "observation_identity", "observation_snapshot",
                "observation_version_identity", "observation_version_snapshot",
            ):
                with self.subTest(table=table):
                    self.assertEqual(
                        connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0],
                        0,
                    )
        finally:
            connection.close()

    def test_import_persists_registries_rights_raw_and_exact_round_trips(self) -> None:
        result = self._import()
        self.assertEqual(len(result.observation_content_hashes), 48)
        self.assertEqual(len(result.observation_version_content_hashes), 48)
        source = self.database.load_source_registry(SOURCE_ID)
        instrument = self.database.load_instrument_registry(INSTRUMENT_ID)
        rights = self.database.load_source_usage_rights(RIGHTS_PROFILE_ID)
        raw = self.database.load_raw_payload(result.raw_payload_hash)
        capture = self.database.load_raw_capture(result.capture_id)
        self.assertEqual(source["access_channel"], ACCESS_CHANNEL)
        self.assertEqual(instrument["source_symbol"], "Copper")
        self.assertEqual(instrument["active_from"], "2022-01-01")
        self.assertEqual(rights["backtest"], "REVIEW_REQUIRED")
        self.assertEqual(rights["redistribution"], "REVIEW_REQUIRED")
        self.assertEqual(rights["reviewed_at"], self.CANONICAL_RIGHTS_REVIEWED_AT)
        self.assertIsNotNone(rights["reviewed_at"])
        self.assertEqual(rights["evidence_reference"], RIGHTS_EVIDENCE_REFERENCE)
        self.assertEqual(raw["byte_size"], len(self.raw))
        self.assertEqual(capture["raw_payload_hash"], result.raw_payload_hash)
        observation = self.database.load_observation(result.observation_content_hashes[0])
        version = self.database.load_observation_version(
            result.observation_version_content_hashes[0]
        )
        self.assertEqual(observation.semantic_data["value"], Decimal("8000.07"))
        self.assertEqual(version.raw_payload_hash, result.raw_payload_hash)
        self.assertEqual(TRANSFORMATION_VERSION, "C4_DB2_WB_COPPER_MONTHLY_V1@1.0.1")
        self.assertEqual(version.transformation_version, TRANSFORMATION_VERSION)
        self.assertEqual(version.semantic_data, observation.semantic_data)
        self.assertIsNone(version.revision_available_at)

    def test_same_raw_reimport_is_idempotent_and_retains_first_metadata(self) -> None:
        first = self._import(suffix="00")
        second = self._import(
            suffix="30",
            rights_reviewed_at="2026-10-02T12:34:56Z",
            source_locator_safe="LATER_EQUIVALENT_LOCATOR",
        )
        self.assertEqual(first, second)
        self.assertEqual(
            self.database.load_source_usage_rights(RIGHTS_PROFILE_ID)["reviewed_at"],
            self.CANONICAL_RIGHTS_REVIEWED_AT,
        )
        self.assertEqual(
            self.database.load_raw_payload(first.raw_payload_hash)["first_seen_at"],
            "2026-10-01T00:00:00Z",
        )
        self.assertEqual(
            self.database.load_raw_capture(first.capture_id)["source_locator_safe"],
            ACCESS_CHANNEL,
        )
        connection = duckdb.connect(str(self.database_path), read_only=True)
        try:
            counts = {
                table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                for table in (
                    "source_registry", "subject_registry", "instrument_registry",
                    "source_usage_rights", "raw_payload", "raw_capture",
                    "observation_identity", "observation_snapshot",
                    "observation_version_identity", "observation_version_snapshot",
                )
            }
        finally:
            connection.close()
        self.assertEqual(counts["raw_payload"], 1)
        self.assertEqual(counts["raw_capture"], 1)
        self.assertEqual(counts["observation_identity"], 48)
        self.assertEqual(counts["observation_snapshot"], 48)
        self.assertEqual(counts["observation_version_identity"], 48)
        self.assertEqual(counts["observation_version_snapshot"], 48)

    def test_rights_evidence_and_decisions_are_exactly_the_approved_profile(self) -> None:
        self._import()
        rights = self.database.load_source_usage_rights(RIGHTS_PROFILE_ID)
        evidence = rights["evidence_reference"]
        for required in (
            "World Bank",
            "0038238",
            "Creative Commons Attribution 4.0",
            "private historical pilot",
        ):
            with self.subTest(required=required):
                self.assertIn(required, evidence)
        for forbidden in ("FRED", "Westmetall", "IMF"):
            with self.subTest(forbidden=forbidden):
                self.assertNotIn(forbidden, evidence)
        self.assertEqual(
            {key: rights[key] for key in (
                "automated_access", "private_storage", "historical_archive",
                "internal_analysis", "backtest", "prediction", "internal_display",
                "internet_display", "redistribution",
            )},
            {
                "automated_access": "ALLOWED",
                "private_storage": "ALLOWED",
                "historical_archive": "ALLOWED",
                "internal_analysis": "ALLOWED",
                "backtest": "REVIEW_REQUIRED",
                "prediction": "REVIEW_REQUIRED",
                "internal_display": "REVIEW_REQUIRED",
                "internet_display": "REVIEW_REQUIRED",
                "redistribution": "REVIEW_REQUIRED",
            },
        )

    def test_different_raw_workbook_retains_old_lineage(self) -> None:
        first = self._import(suffix="00")
        changed = workbook_bytes(value_overrides={"2024-01": 9999.25})
        second = self._import(changed, suffix="30")
        self.assertNotEqual(first.raw_payload_hash, second.raw_payload_hash)
        connection = duckdb.connect(str(self.database_path), read_only=True)
        try:
            raw_count = connection.execute("SELECT COUNT(*) FROM raw_payload").fetchone()[0]
            version_count = connection.execute(
                "SELECT COUNT(*) FROM observation_version_identity"
            ).fetchone()[0]
            old_count = connection.execute(
                "SELECT COUNT(*) FROM observation_version_identity WHERE raw_payload_hash = ?",
                [first.raw_payload_hash],
            ).fetchone()[0]
        finally:
            connection.close()
        self.assertEqual(raw_count, 2)
        self.assertEqual(version_count, 96)
        self.assertEqual(old_count, 48)
        self.database.load_observation_version(first.observation_version_content_hashes[0])

    def test_real_storage_does_not_enable_readiness_or_operational_paths(self) -> None:
        self._import()
        self.assertTrue(all(value is False for value in SAFETY_FLAGS.values()))
        self.assertEqual(PIT_ENABLED_SOURCES, ())
        self.assertEqual(RESEARCH_ENABLED_SOURCES, ())
        self.assertEqual(BACKTEST_ENABLED_SOURCES, ())
        self.assertEqual(len(RD3_OPEN_BLOCKERS), 15)


if __name__ == "__main__":
    unittest.main()
